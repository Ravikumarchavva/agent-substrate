"""Regenerates ``tests/fixtures/documents/`` — real files, made by real office software, so the reader is tested on what producers write.

    uv run python tests/fixtures/make_documents.py

The DOCX/PPTX/XLSX are written with python-docx / python-pptx / openpyxl; LibreOffice (headless) then recalculates the workbook (so
formula cells carry cached values, as Excel writes them) and converts each to its OpenDocument twin. The files are committed:
the tests do not need any of these tools."""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

OUT = Path(__file__).parent / "documents"
WORK = Path(__file__).parent / "_work"  # inside the repo: a snap-packaged LibreOffice cannot see /tmp


def docx(path: Path) -> None:
    import docx as D
    from docx.enum.text import WD_BREAK

    d = D.Document()
    d.core_properties.title = "Quarterly Operations Report"
    d.add_heading("Quarterly Operations Report", 0)
    d.add_paragraph("Prepared for the board. Invoice 4417 was settled in full.")
    d.add_heading("Revenue", 1)
    d.add_paragraph("Revenue grew while costs stayed flat.")
    d.add_paragraph("Direct sales", style="List Bullet")
    d.add_paragraph("Partner channel", style="List Bullet")
    d.add_paragraph("First step", style="List Number")
    d.add_paragraph("Second step", style="List Number")
    d.add_heading("Regional detail", 2)
    table = d.add_table(rows=3, cols=3)
    for r, row in enumerate([("Region", "Q2", "Q3"), ("EMEA", "120", "140"), ("APAC", "80", "95")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    p = d.add_paragraph("Details are on the next page. See ")
    p.add_run("the dashboard").font.underline = True
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("Risks", 1)
    d.add_paragraph("Shipping to Rotterdam may be delayed.")
    d.save(path)


def pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Launch Plan"
    s.placeholders[1].text = "Autumn campaign"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Milestones"
    body = s.placeholders[1].text_frame
    body.text = "Design freeze"
    body.add_paragraph().text = "Beta release"
    sub = body.add_paragraph()
    sub.text = "Invite 50 customers"
    sub.level = 1
    s.notes_slide.notes_text_frame.text = "Mention the Rotterdam warehouse."
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Budget"
    rows = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for r, (a, b) in enumerate([("Item", "EUR"), ("Ads", "12000"), ("Events", "8000")]):
        rows.cell(r, 0).text, rows.cell(r, 1).text = a, b
    prs.save(path)


def xlsx(path: Path) -> None:
    import datetime as dt

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Units", "Price", "Revenue", "Closed on"])
    ws.append(["Jan", 10, 2.5, "=B2*C2", dt.date(2026, 1, 31)])
    ws.append(["Feb", 12, 2.5, "=B3*C3", dt.date(2026, 2, 28)])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Remark"])
    ws2.append(["Totals exclude returns | pipes and\nnewlines"])
    wb.create_sheet("Helper").sheet_state = "hidden"
    wb["Helper"].append(["do not show"])
    wb.save(path)


def convert(src: Path, fmt: str, outdir: Path) -> Path:
    profile = WORK / "profile"
    subprocess.run(
        [shutil.which("soffice") or shutil.which("libreoffice") or "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", fmt, "--outdir", str(outdir), str(src)],
        check=True,
        capture_output=True,
        timeout=180,
    )
    return outdir / (src.stem + "." + fmt.split(":")[0])


def pdfs() -> None:
    """``scanned_page.pdf`` (page 1 is a picture of text with no text layer; page 2 has a real text layer) and ``bookmarks.pdf``."""
    sys.path.insert(0, str(Path(__file__).parent))
    from pdfs import Line, Page, build, scan_of  # noqa: PLC0415

    scan = scan_of([Line("Invoice 4417 total due 120 EUR", 18), Line("Shipping to Rotterdam on Friday", 18, gap=10)])
    (OUT / "scanned_page.pdf").write_bytes(build([Page(image=scan), Page([Line("Second page has a text layer", 14)])]))
    body = [Line("Body text sits here and runs on for a while so that the body size is clear.", 11)]
    pages = [
        Page([Line("Chapter One", 24, bold=True), Line("Intro paragraph of chapter one.", 11, gap=8), *body]),
        Page([Line("Background", 16, bold=True), *body, Line("Details", 16, bold=True, gap=14), *body]),
    ]
    (OUT / "bookmarks.pdf").write_bytes(
        build(pages, outline=[("Chapter One", 1, 0), ("Background", 2, 1), ("Details", 2, 1)], title="Chapter Book")
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.rmtree(WORK, ignore_errors=True)
    WORK.mkdir()
    work = WORK
    docx(work / "sample.docx")
    pptx(work / "sample.pptx")
    xlsx(work / "raw.xlsx")
    recalculated = work / "recalc"
    recalculated.mkdir()
    shutil.copy(convert(work / "raw.xlsx", "xlsx", recalculated), OUT / "sample.xlsx")
    shutil.copy(work / "sample.docx", OUT / "sample.docx")
    shutil.copy(work / "sample.pptx", OUT / "sample.pptx")
    for src, fmt in [(OUT / "sample.docx", "odt"), (OUT / "sample.pptx", "odp"), (OUT / "sample.xlsx", "ods")]:
        shutil.copy(convert(src, fmt, work), OUT / f"sample.{fmt}")
    pdfs()
    shutil.rmtree(WORK, ignore_errors=True)
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:14} {f.stat().st_size:7} bytes")


if __name__ == "__main__":
    sys.exit(main())
